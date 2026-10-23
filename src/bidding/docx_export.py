"""Convert the generated Markdown bid into a Word document."""

from __future__ import annotations

import os
import tempfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from markdown_it import MarkdownIt
from markdown_it.tree import SyntaxTreeNode

from .errors import BidExportError

BODY_FONT_EAST_ASIA = "宋体"
BODY_FONT_ASCII = "Times New Roman"
HEADING_FONT_EAST_ASIA = "黑体"


@dataclass(frozen=True, slots=True)
class DocxResult:
    path: Path
    paragraphs: int
    tables: int


def _set_run_font(run, *, east_asia: str, ascii_font: str, size: float | None = None) -> None:
    run.font.name = ascii_font
    if size is not None:
        run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)


def _set_style_font(style, *, east_asia: str, ascii_font: str, size: float | None = None) -> None:
    style.font.name = ascii_font
    if size is not None:
        style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)


def _plain(node: SyntaxTreeNode) -> str:
    if node.type == "text":
        return node.content
    if node.type == "code_inline":
        return node.content
    if node.type in ("softbreak", "hardbreak"):
        return "\n"
    if node.type == "image":
        return f"【图片：{node.attrs.get('alt', '')}】"
    return "".join(_plain(child) for child in node.children)


def _add_inline(
    paragraph,
    node: SyntaxTreeNode,
    *,
    bold: bool = False,
    italic: bool = False,
) -> None:
    if node.type == "text" or node.type == "code_inline":
        run = paragraph.add_run(node.content)
        run.bold = bold or None
        run.italic = italic or None
        if node.type == "code_inline":
            _set_run_font(run, east_asia=BODY_FONT_EAST_ASIA, ascii_font="Consolas")
        else:
            _set_run_font(run, east_asia=BODY_FONT_EAST_ASIA, ascii_font=BODY_FONT_ASCII)
        return
    if node.type in ("softbreak", "hardbreak"):
        paragraph.add_run("\n" if node.type == "hardbreak" else " ")
        return
    if node.type == "image":
        run = paragraph.add_run(f"【图片：{node.attrs.get('alt', '')}】")
        _set_run_font(run, east_asia=BODY_FONT_EAST_ASIA, ascii_font=BODY_FONT_ASCII)
        return
    child_bold = bold or node.type == "strong"
    child_italic = italic or node.type == "em"
    for child in node.children:
        _add_inline(paragraph, child, bold=child_bold, italic=child_italic)


def _list_style(ordered: bool, depth: int) -> str:
    base = "List Number" if ordered else "List Bullet"
    if depth <= 1:
        return base
    return f"{base} {min(depth, 3)}"


def _safe_style(document, style: str, fallback: str) -> str:
    try:
        document.styles[style]
        return style
    except KeyError:
        return fallback


def _render_list(node: SyntaxTreeNode, document, depth: int = 1) -> None:
    ordered = node.type == "ordered_list"
    style = _safe_style(
        document,
        _list_style(ordered, depth),
        "List Number" if ordered else "List Bullet",
    )
    for item in node.children:
        if item.type != "list_item":
            continue
        paragraphs = [child for child in item.children if child.type == "paragraph"]
        nested = [
            child for child in item.children
            if child.type in ("bullet_list", "ordered_list")
        ]
        if paragraphs:
            paragraph = document.add_paragraph(style=style)
            for child in paragraphs[0].children:
                _add_inline(paragraph, child)
            for extra in paragraphs[1:]:
                paragraph = document.add_paragraph(style=style)
                for child in extra.children:
                    _add_inline(paragraph, child)
        for child in nested:
            _render_list(child, document, depth + 1)


def _table_rows(node: SyntaxTreeNode) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    header_rows = 0
    for section in node.children:
        if section.type not in ("thead", "tbody"):
            continue
        for row in section.children:
            if row.type != "tr":
                continue
            cells = [_plain(cell).strip() for cell in row.children if cell.type in ("th", "td")]
            if cells:
                rows.append(cells)
                if section.type == "thead":
                    header_rows += 1
    return rows, header_rows


def _render_table(node: SyntaxTreeNode, document) -> None:
    rows, header_rows = _table_rows(node)
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column, value in enumerate(row):
            cell = table.cell(row_index, column)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(value)
            _set_run_font(run, east_asia=BODY_FONT_EAST_ASIA, ascii_font=BODY_FONT_ASCII, size=10.5)
            run.bold = row_index < header_rows or None


def _configure_document(document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    normal = document.styles["Normal"]
    _set_style_font(normal, east_asia=BODY_FONT_EAST_ASIA, ascii_font=BODY_FONT_ASCII, size=10.5)
    normal.paragraph_format.line_spacing = 1.5

    heading_sizes = {"Heading 1": 18, "Heading 2": 16, "Heading 3": 14, "Heading 4": 12}
    for name, size in heading_sizes.items():
        style = document.styles[name]
        _set_style_font(style, east_asia=HEADING_FONT_EAST_ASIA, ascii_font=HEADING_FONT_EAST_ASIA, size=size)


def markdown_to_docx(
    markdown: str,
    output_path: str | Path,
    *,
    title: str | None = None,
) -> DocxResult:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    try:
        document = Document()
        _configure_document(document)
        if title:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(title)
            run.bold = True
            _set_run_font(run, east_asia=HEADING_FONT_EAST_ASIA, ascii_font=HEADING_FONT_EAST_ASIA, size=18)

        parser = MarkdownIt("commonmark").enable("table").enable("strikethrough")
        root = SyntaxTreeNode(parser.parse(markdown or ""))
        for node in root.children:
            if node.type == "heading":
                level = int(node.tag[1]) if node.tag and node.tag[1:].isdigit() else 1
                document.add_heading(_plain(node).strip(), level=min(max(level, 1), 4))
            elif node.type == "paragraph":
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.first_line_indent = Pt(21)
                for child in node.children:
                    _add_inline(paragraph, child)
            elif node.type in ("bullet_list", "ordered_list"):
                _render_list(node, document)
            elif node.type == "table":
                _render_table(node, document)
            elif node.type in ("fence", "code_block"):
                paragraph = document.add_paragraph()
                run = paragraph.add_run(node.content.rstrip("\n"))
                _set_run_font(run, east_asia=BODY_FONT_EAST_ASIA, ascii_font="Consolas", size=9)
            elif node.type == "blockquote":
                for child in node.children:
                    if child.type == "paragraph":
                        paragraph = document.add_paragraph(
                            style=_safe_style(document, "Quote", "Normal")
                        )
                        for inline in child.children:
                            _add_inline(paragraph, inline)
            elif node.type == "hr":
                document.add_paragraph("─" * 30)

        fd, temp_name = tempfile.mkstemp(prefix=".bid-", suffix=".docx", dir=str(output.parent))
        os.close(fd)
        try:
            document.save(temp_name)
            os.replace(temp_name, output)
        except BaseException:
            try:
                os.unlink(temp_name)
            except OSError:
                pass
            raise
        return DocxResult(
            path=output,
            paragraphs=len(document.paragraphs),
            tables=len(document.tables),
        )
    except BidExportError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise BidExportError(f"failed to export Word document: {exc}") from exc
