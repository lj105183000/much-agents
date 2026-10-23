from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from src.bidding.docx_export import markdown_to_docx
from src.bidding.errors import BidInputError
from src.bidding.profile import BidderProfile
from src.bidding.runner import run_bid_writing
from tests.fakes import (
    FakeChatClient,
    FakeRetriever,
    make_bid_settings,
    write_parsed_fixture,
)


def test_runner_writes_complete_artifacts(tmp_path: Path) -> None:
    write_parsed_fixture(tmp_path)
    settings = make_bid_settings(tmp_path)
    result = run_bid_writing(
        "doc-1",
        profile=BidderProfile.empty(),
        max_sections=2,
        settings=settings,
        run_id="run-test",
        llm=FakeChatClient(),
        retriever=FakeRetriever(),
    )
    assert result.ok is True
    assert result.exit_code == 0
    assert result.status == "degraded"
    assert result.narrative_written == 2
    for path in (
        result.bid_markdown,
        result.bid_docx,
        result.report_path,
        result.outline_path,
        result.output_dir / "profile.json",
    ):
        assert Path(path).is_file()
    assert len(list(result.sections_dir.glob("*.md"))) == 11

    markdown = result.bid_markdown.read_text(encoding="utf-8")
    assert "# 测试供水工程" in markdown
    assert "## 一、投标函及投标函附录" in markdown
    assert "### 6.1 编制说明与工程概况" in markdown
    assert "### 6.2 施工方案与技术措施" in markdown
    assert "### 6.3 质量管理体系与措施" in markdown
    assert "【待填写：公司名称】" in markdown
    assert "【待填写：投标总价" in markdown
    assert "999万元" not in markdown

    document = Document(str(result.bid_docx))
    assert len(document.paragraphs) > 20
    assert len(document.tables) >= 5
    assert any("测试供水工程" in paragraph.text for paragraph in document.paragraphs)

    import json

    report = json.loads(result.report_path.read_text(encoding="utf-8"))
    assert report["kind"] == "bid"
    assert report["coverage"]["total"] >= 5
    assert report["coverage"]["counts"]["template"] >= 1
    assert report["missing_profile_fields"]
    assert report["placeholders"]["count"] > 0


def test_runner_rejects_existing_run_artifacts(tmp_path: Path) -> None:
    write_parsed_fixture(tmp_path)
    settings = make_bid_settings(tmp_path)
    kwargs = dict(
        profile=BidderProfile.empty(),
        max_sections=1,
        settings=settings,
        run_id="same-run",
        llm=FakeChatClient(),
        retriever=FakeRetriever(),
    )
    run_bid_writing("doc-1", **kwargs)  # type: ignore[arg-type]
    with pytest.raises(BidInputError):
        run_bid_writing("doc-1", **kwargs)  # type: ignore[arg-type]


def test_runner_without_profile_uses_placeholders(tmp_path: Path) -> None:
    write_parsed_fixture(tmp_path)
    settings = make_bid_settings(tmp_path)
    result = run_bid_writing(
        "doc-1",
        max_sections=1,
        settings=settings,
        run_id="no-profile",
        llm=FakeChatClient(),
        retriever=FakeRetriever(),
    )
    assert any("no bidder profile" in warning for warning in result.warnings)
    markdown = result.bid_markdown.read_text(encoding="utf-8")
    assert "【待填写：公司名称】" in markdown


def test_markdown_to_docx_headings_lists_and_tables(tmp_path: Path) -> None:
    markdown = """# 标题

正文 **加粗** 与 *斜体*。

- 条目一
- 条目二

| 序号 | 名称 |
| --- | --- |
| 1 | 测试 |
"""
    output = tmp_path / "out.docx"
    result = markdown_to_docx(markdown, output, title="测试文档")
    assert result.path.is_file()
    document = Document(str(output))
    assert any(paragraph.style.name.startswith("Heading") for paragraph in document.paragraphs)
    assert len(document.tables) == 1
    assert document.tables[0].cell(1, 1).text == "测试"
