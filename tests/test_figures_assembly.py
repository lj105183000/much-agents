from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document
from PIL import Image

from src.bidding.assemble import assemble_bid
from src.bidding.checks import build_allowed
from src.bidding.docx_export import markdown_to_docx
from src.bidding.errors import BidExportError
from src.bidding.figures import build_figure_blocks, generate_figures
from src.bidding.gate import evaluate_gate
from src.bidding.plan import plan_bid
from src.bidding.profile import BidderProfile
from src.bidding.staged_output import SectionOutput, assemble_staged_bid, write_staged_outputs
from tests.fakes import make_bid_settings, write_parsed_fixture

PROFILE = {
    "company": {"name": "测试建设有限公司", "address": "北京市", "phone": "010-1"},
    "project_manager": {"name": "王五"},
    "personnel": [{"role": "技术负责人", "name": "赵六"}],
}


def _fixture(tmp_path: Path):
    write_parsed_fixture(tmp_path)
    settings = make_bid_settings(
        tmp_path,
        BID_FIGURES_ENABLED="true",
        BID_FIGURES_USE_LLM="false",
    )
    planned = plan_bid("doc-1", settings, max_sections=2)
    profile = BidderProfile.from_mapping(PROFILE)
    first, second = planned.outline.selected_narratives
    outputs = {
        first.id: SectionOutput(
            section_id=first.id,
            status="final",
            final_markdown="项目名称：测试供水工程\n",
        ),
        second.id: SectionOutput(
            section_id=second.id,
            status="final",
            final_markdown="1. 施工准备\n2. 测量放线\n3. 土方开挖\n4. 基础施工\n",
        ),
    }
    allowed = build_allowed(planned.context, profile, planned.outline)
    base = assemble_staged_bid(
        context=planned.context,
        profile=profile,
        outline=planned.outline,
        outputs=outputs,
        allowed=allowed,
    )
    manifest = generate_figures(
        context=planned.context,
        profile=profile,
        outline=planned.outline,
        section_markdowns=base.section_markdowns,
        outputs=outputs,
        run_dir=tmp_path / "run",
        settings=settings,
    )
    return planned, profile, outputs, allowed, manifest, settings


def test_assembly_inserts_figures_and_replaces_placeholders(tmp_path: Path) -> None:
    planned, profile, outputs, allowed, manifest, _settings = _fixture(tmp_path)
    assembly = assemble_staged_bid(
        context=planned.context,
        profile=profile,
        outline=planned.outline,
        outputs=outputs,
        allowed=allowed,
        figure_blocks=build_figure_blocks(manifest),
        figure_stats=manifest.stats,
        figure_warnings=manifest.warnings,
        figures_enabled=manifest.enabled,
    )
    assert assembly.markdown.count("](figures/") == 4
    assert "【待填写：施工进度网络图/横道图】" not in assembly.section_markdowns[
        "att-06-04"
    ]
    assert "【待填写：施工总平面图】" not in assembly.section_markdowns["att-06-05"]
    assert assembly.report["figures"]["stats"]["generated"] == 4
    assert assembly.report["stats"]["figures"] == 4
    assert assembly.report["stats"]["estimated_pages"] >= 2
    gate = evaluate_gate(
        outline=planned.outline,
        outputs=outputs,
        report=assembly.report,
        missing_profile_fields=[],
    )
    assert "figures" in [check["id"] for check in gate["checks"]]


def test_figure_block_appends_when_placeholder_is_missing(tmp_path: Path) -> None:
    from src.bidding.figures.schema import FigureBlock

    planned, profile, outputs, allowed, _manifest, _settings = _fixture(tmp_path)
    block = FigureBlock(
        section_id="att-06-05",
        markdown="![图 9 补充图](figures/x.png)\n\n*图 9 补充图*",
        replace_text="不存在的占位符",
    )
    assembly = assemble_bid(
        context=planned.context,
        profile=profile,
        outline=planned.outline,
        drafts=[],
        reviews=[],
        revisions=[],
        warnings=[],
        allowed=allowed,
        figure_blocks=[block],
    )
    assert assembly.section_markdowns["att-06-05"].rstrip().endswith("*图 9 补充图*")


def test_write_staged_outputs_embeds_images_and_fixes_section_paths(
    tmp_path: Path,
) -> None:
    planned, profile, outputs, allowed, manifest, settings = _fixture(tmp_path)
    assembly = assemble_staged_bid(
        context=planned.context,
        profile=profile,
        outline=planned.outline,
        outputs=outputs,
        allowed=allowed,
        figure_blocks=build_figure_blocks(manifest),
        figure_stats=manifest.stats,
        figure_warnings=manifest.warnings,
    )
    result = write_staged_outputs(
        context=planned.context,
        assembly=assembly,
        outline=planned.outline,
        output_dir=tmp_path / "run",
        image_width_cm=settings.figure_image_width_cm,
    )
    assert result["docx_images"] == 4
    with zipfile.ZipFile(result["bid_docx"]) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert len(media) == 4
    section_text = (
        tmp_path / "run" / "sections" / f"{planned.outline.selected_narratives[1].id}.md"
    ).read_text(encoding="utf-8")
    assert "](../figures/" in section_text
    report = result["report"]
    assert report["docx_images"] == 4
    assert report["figures"]["docx_images"] == 4


def _make_png(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (120, 80), "white").save(path)


def test_docx_embeds_safe_image_and_preserves_legacy_placeholder(tmp_path: Path) -> None:
    _make_png(tmp_path / "figures" / "a.png")
    markdown = "![图 1](figures/a.png)\n\n图 1 说明"
    output = tmp_path / "out.docx"
    result = markdown_to_docx(markdown, output, asset_dir=tmp_path)
    assert result.images == 1
    with zipfile.ZipFile(output) as archive:
        assert len([n for n in archive.namelist() if n.startswith("word/media/")]) == 1
    document = Document(str(output))
    centered = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.alignment is not None and paragraph.alignment.name == "CENTER"
    ]
    assert centered

    legacy = markdown_to_docx(markdown, tmp_path / "legacy.docx")
    assert legacy.images == 0
    assert "【图片：图 1】" in Document(str(legacy.path)).paragraphs[0].text


@pytest.mark.parametrize(
    "source",
    [
        "../secret.png",
        "C:/secret.png",
        "figures\\a.png",
        "https://example.com/a.png",
        "figures/missing.png",
        "figures/a.svg",
    ],
)
def test_docx_rejects_unsafe_or_missing_images(tmp_path: Path, source: str) -> None:
    _make_png(tmp_path / "figures" / "a.png")
    with pytest.raises(BidExportError):
        markdown_to_docx(
            f"![图]({source})",
            tmp_path / "bad.docx",
            asset_dir=tmp_path,
        )


def test_gate_warns_when_figures_were_skipped_or_failed(tmp_path: Path) -> None:
    planned, profile, outputs, allowed, manifest, _settings = _fixture(tmp_path)
    assembly = assemble_staged_bid(
        context=planned.context,
        profile=profile,
        outline=planned.outline,
        outputs=outputs,
        allowed=allowed,
        figure_blocks=build_figure_blocks(manifest),
        figure_stats={**manifest.stats, "skipped": 2, "failed": 1},
    )
    gate = evaluate_gate(
        outline=planned.outline,
        outputs=outputs,
        report=assembly.report,
        missing_profile_fields=[],
    )
    check = next(item for item in gate["checks"] if item["id"] == "figures")
    assert check["level"] == "warn"
    assert check["status"] == "fail"
    assert "1" in check["message"] and "2" in check["message"]
